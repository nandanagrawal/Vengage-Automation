import csv
import re
import sys
from datetime import datetime

sys.path.insert(0, "/home/dell/Documents/Vengage-Automation/backend")

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.db.session import SessionLocal
from app.models.customer import Customer
from app.models.customer_product_and_service import CustomerProductAndService
from app.models.product_and_service import ProductAndService

CSV_PATH = "/home/dell/Downloads/Customer-Services-And-Billing-Master-Data.csv"
OUT_PATH = "/home/dell/Downloads/Customer-Product-Gaps-Analysis.xlsx"

NAVY = "1B4F72"
LIGHT_BLUE = "D6E4F0"
WHITE = "FFFFFF"
ALT_ROW = "F2F7FC"
WARN_BG = "FDECEA"
WARN_FG = "B3261E"
BORDER_CLR = "C5D3DF"
THIN = Side(style="thin", color=BORDER_CLR)
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(name="Calibri", bold=True, color=WHITE, size=10)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BODY_FONT = Font(name="Calibri", size=10)
ALT_FILL = PatternFill("solid", fgColor=ALT_ROW)


def norm_customer(name: str) -> str:
    n = name.strip().lower()
    n = re.sub(r"^\([^)]*\)\s*", "", n)
    n = re.sub(r"[^a-z0-9]+", "", n)
    return n


def norm_product(name: str) -> str:
    first_line = name.split("\n", 1)[0]
    n = first_line.strip().lower()
    n = re.sub(r"\s+", " ", n)
    return n


def add_sheet(wb, title, headers, rows, note=None):
    ws = wb.create_sheet(title)
    start_row = 1
    if note:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        c = ws.cell(row=1, column=1, value=note)
        c.font = Font(name="Calibri", italic=True, size=9, color="5A6B7A")
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28
        start_row = 2

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_i, row in enumerate(rows, start=start_row + 1):
        is_alt = (r_i - start_row) % 2 == 0
        for c_i, val in enumerate(row, start=1):
            cell = ws.cell(row=r_i, column=c_i, value=val)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False, indent=1 if c_i == 1 else 0)
            if is_alt:
                cell.fill = ALT_FILL

    for col_idx, h in enumerate(headers, start=1):
        max_len = max([len(str(h))] + [len(str(row[col_idx - 1])) for row in rows] or [10])
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, min(60, max_len * 0.95 + 2))

    last_row = start_row + len(rows)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{last_row}"
    ws.freeze_panes = f"A{start_row + 1}"
    return ws


# ── Parse CSV ─────────────────────────────────────────────────────────────
with open(CSV_PATH, newline="", encoding="utf-8-sig") as f:
    reader = csv.reader(f)
    header = next(reader)
    service_cols = header[2:]
    csv_customers: dict[str, list[str]] = {}
    for row in reader:
        if not row or not row[0].strip():
            continue
        name = row[0].strip()
        marked = [col for i, col in enumerate(service_cols, start=2) if i < len(row) and row[i].strip()]
        csv_customers[name] = marked

# ── DB state ──────────────────────────────────────────────────────────────
db = SessionLocal()
try:
    customers = db.query(Customer).all()
    by_norm_cust: dict[str, list[Customer]] = {}
    for c in customers:
        by_norm_cust.setdefault(norm_customer(c.display_name), []).append(c)

    products = db.query(ProductAndService).all()
    by_norm_prod: dict[str, list[ProductAndService]] = {}
    for p in products:
        by_norm_prod.setdefault(norm_product(p.name), []).append(p)

    cps_rows = db.query(CustomerProductAndService).all()
    mapped_pairs = {(c.customer_id, c.product_and_service_id) for c in cps_rows}
    cps_by_customer: dict[int, list[CustomerProductAndService]] = {}
    for cps in cps_rows:
        cps_by_customer.setdefault(cps.customer_id, []).append(cps)
    products_by_id = {p.id: p for p in products}
    cust_by_id = {c.id: c for c in customers}

    def resolve_customer(cust_name: str):
        key = norm_customer(cust_name)
        matches = by_norm_cust.get(key)
        if matches:
            return matches[0]
        # Fallback: CSV name carries a trailing qualifier the app's name doesn't
        # (e.g. "... - Different Customers") — try a prefix match either way.
        if len(key) < 10:
            return None
        for db_key, db_matches in by_norm_cust.items():
            if len(db_key) < 10:
                continue
            if key.startswith(db_key) or db_key.startswith(key):
                return db_matches[0]
        return None

    all_marked_headers = sorted({h for marks in csv_customers.values() for h in marks})
    header_to_product: dict[str, ProductAndService | None] = {}
    for h in all_marked_headers:
        matches = by_norm_prod.get(norm_product(h))
        header_to_product[h] = matches[0] if matches else None

    # 1. Customers in the sheet not found in the app at all
    missing_customers_rows = []
    for cust_name, marks in csv_customers.items():
        if not marks:
            continue
        if resolve_customer(cust_name) is None:
            missing_customers_rows.append([
                cust_name,
                ", ".join(h.split("\n", 1)[0] for h in marks),
            ])

    # 2. Customers found in the app but with zero services mapped
    zero_service_rows = []
    for cust_name, marks in csv_customers.items():
        if not marks:
            continue
        cust = resolve_customer(cust_name)
        if cust is None:
            continue
        if not cps_by_customer.get(cust.id):
            zero_service_rows.append([
                cust_name, cust.display_name,
                ", ".join(h.split("\n", 1)[0] for h in marks),
            ])

    # 3. Customer x product gaps
    product_gap_rows = []
    for cust_name, marks in csv_customers.items():
        cust = resolve_customer(cust_name)
        if cust is None:
            continue
        for h in marks:
            ps = header_to_product.get(h)
            if ps is None:
                continue
            if (cust.id, ps.id) not in mapped_pairs:
                product_gap_rows.append([cust_name, cust.display_name, ps.name])

    # 4. Sheet products with no match in the app's catalog at all
    unresolved_rows = []
    for h in all_marked_headers:
        if header_to_product.get(h) is None:
            n = sum(1 for marks in csv_customers.values() if h in marks)
            unresolved_rows.append([h.replace("\n", " / "), n])

    # 5. Mapped, but sheet_column_id missing -> produces no invoice line item
    no_column_rows = []
    for cps in cps_rows:
        if cps.sheet_column_id is None:
            cust = cust_by_id.get(cps.customer_id)
            ps = products_by_id.get(cps.product_and_service_id)
            no_column_rows.append([
                cust.display_name if cust else f"customer#{cps.customer_id}",
                ps.name if ps else f"product#{cps.product_and_service_id}",
            ])
    no_column_rows.sort(key=lambda r: r[0])

    # ── Build workbook ────────────────────────────────────────────────────
    wb = Workbook()
    wb.remove(wb.active)

    ws0 = wb.create_sheet("Summary")
    ws0.merge_cells("A1:B1")
    t = ws0["A1"]
    t.value = f"Customer/Product Gap Analysis — vs. master sheet — generated {datetime.now().strftime('%d %b %Y, %H:%M')}"
    t.font = Font(name="Calibri", bold=True, size=13, color=WHITE)
    t.fill = PatternFill("solid", fgColor="0D3349")
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws0.row_dimensions[1].height = 26
    ws0.column_dimensions["A"].width = 46
    ws0.column_dimensions["B"].width = 12
    summary = [
        ("Customers in sheet missing from the app", len(missing_customers_rows)),
        ("Customers with zero services mapped", len(zero_service_rows)),
        ("Customer x product gaps (existing customer/product, not linked)", len(product_gap_rows)),
        ("Sheet products not found in the app's catalog at all", len(unresolved_rows)),
        ("Mapped services with no sheet column set", len(no_column_rows)),
    ]
    for i, (label, val) in enumerate(summary, start=3):
        lc = ws0.cell(row=i, column=1, value=label)
        vc = ws0.cell(row=i, column=2, value=val)
        lc.font = Font(name="Calibri", size=11, bold=True)
        vc.font = Font(name="Calibri", size=11, bold=True, color=WARN_FG if val else "1A6B2A")
        lc.border = THIN_BORDER
        vc.border = THIN_BORDER
        vc.alignment = Alignment(horizontal="center")
        if (i - 3) % 2 == 0:
            lc.fill = ALT_FILL
            vc.fill = ALT_FILL

    add_sheet(
        wb, "Missing Customers",
        ["Customer name (per sheet)", "Products marked for them in the sheet"],
        missing_customers_rows,
        note="Customers the master sheet marks as having services, but who don't exist in the app at all yet.",
    )
    add_sheet(
        wb, "Zero Services Mapped",
        ["Customer name (per sheet)", "Customer name (app)", "Products marked for them in the sheet"],
        zero_service_rows,
        note="Customer exists in the app, but has no product/service mapped to them at all.",
    )
    add_sheet(
        wb, "Customer x Product Gaps",
        ["Customer name (per sheet)", "Customer name (app)", "Missing product"],
        product_gap_rows,
        note="Product exists in the app's catalog, customer exists in the app, but they're not linked.",
    )
    add_sheet(
        wb, "Products Not In Catalog",
        ["Product (per sheet header)", "# customers marked for it"],
        unresolved_rows,
        note="No product in the app matches this sheet column header at all — likely not synced from QBO yet.",
    )
    add_sheet(
        wb, "No Sheet Column Set",
        ["Customer name (app)", "Product"],
        no_column_rows,
        note="Product IS linked to the customer, but has no sheet column set — produces no invoice line item yet.",
    )

    wb.save(OUT_PATH)
    print(f"Saved -> {OUT_PATH}")
finally:
    db.close()
